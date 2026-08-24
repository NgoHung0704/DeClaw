"""DOCX via python-docx.

Order is the subtle part. ``document.paragraphs`` and ``document.tables`` are
separate lists, so consuming them in turn silently moves every table to the end
of the document — and a chunk's surrounding context is exactly what makes a
citation trustworthy. Walking ``document.element.body`` and dispatching on tag
name keeps the real order.

Headings are metadata, not content: they become the ``heading`` of the blocks
that follow, the same convention the Markdown parser uses.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from docx import Document
from docx.table import Table
from docx.text.paragraph import Paragraph
from models import Block, ParsedDocument


def _row_text(table: Table) -> list[str]:
    rows: list[str] = []
    for row in table.rows:
        cells = [cell.text.strip() for cell in row.cells]
        if any(cells):
            rows.append(" | ".join(cells))
    return rows


def parse_docx(path: Path) -> ParsedDocument:
    """Extract paragraphs and tables in document order."""
    try:
        document = Document(str(path))
    except Exception as exc:  # noqa: BLE001 - python-docx raises several types
        raise ValueError(f"Cannot read DOCX {path.name}: {exc}") from exc

    blocks: list[Block] = []
    heading: str | None = None
    body: Any = document.element.body

    for child in body.iterchildren():
        tag = str(child.tag).rsplit("}", 1)[-1]
        if tag == "p":
            paragraph = Paragraph(child, document)
            text = paragraph.text.strip()
            if not text:
                continue
            if (paragraph.style.name or "").startswith("Heading"):
                heading = text
                continue
            blocks.append(Block(text=text, heading=heading))
        elif tag == "tbl":
            for row in _row_text(Table(child, document)):
                blocks.append(Block(text=row, heading=heading))

    return ParsedDocument(doc_type="docx", page_count=1, blocks=blocks)
