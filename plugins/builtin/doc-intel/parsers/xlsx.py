"""XLSX via openpyxl, one block per data row.

Each row is rendered with its column headers inline ("Poste: Conseil |
Montant: 12000"). A bare row of values retrieves badly and answers nothing —
the number needs its column name to mean anything. The few extra tokens buy a
chunk that stands on its own.

``read_only=True`` and ``data_only=True``: we want values, not formulas, and we
never want to hold a large workbook fully in memory.
"""

from __future__ import annotations

from pathlib import Path
from typing import Any

from models import Block, ParsedDocument
from openpyxl import load_workbook


def _cell(value: Any) -> str:
    return "" if value is None else str(value).strip()


def parse_xlsx(path: Path) -> ParsedDocument:
    """Extract each sheet's data rows, headers folded into every row."""
    try:
        workbook = load_workbook(str(path), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001 - openpyxl raises several types
        raise ValueError(f"Cannot read XLSX {path.name}: {exc}") from exc

    blocks: list[Block] = []
    try:
        for sheet in workbook.worksheets:
            headers: list[str] = []
            for row in sheet.iter_rows(values_only=True):
                cells = [_cell(value) for value in row]
                if not any(cells):
                    continue
                if not headers:
                    headers = cells
                    continue
                pairs = [
                    f"{headers[i]}: {cell}" if i < len(headers) and headers[i] else cell
                    for i, cell in enumerate(cells)
                    if cell
                ]
                if pairs:
                    blocks.append(Block(text=" | ".join(pairs), sheet=sheet.title))
    finally:
        workbook.close()

    return ParsedDocument(doc_type="xlsx", page_count=1, blocks=blocks)
