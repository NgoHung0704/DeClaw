"""DOCX parsing: headings, order, and tables."""

from __future__ import annotations

from pathlib import Path

import pytest
from docx import Document

from parsers.docx import parse_docx


def _make_docx(path: Path) -> Path:
    document = Document()
    document.add_heading("Contrat de prestation", level=1)
    document.add_paragraph("Preambule du contrat.")
    document.add_heading("Article 3", level=2)
    document.add_paragraph("Obligations du prestataire.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Poste"
    table.cell(0, 1).text = "Montant"
    table.cell(1, 0).text = "Conseil"
    table.cell(1, 1).text = "12000"
    document.add_paragraph("Paragraphe final.")
    document.save(str(path))
    return path


def test_doc_type_and_paragraph_text(tmp_path: Path) -> None:
    parsed = parse_docx(_make_docx(tmp_path / "c.docx"))
    assert parsed.doc_type == "docx"
    texts = [b.text for b in parsed.blocks]
    assert "Preambule du contrat." in texts
    assert "Obligations du prestataire." in texts


def test_headings_become_metadata_not_content(tmp_path: Path) -> None:
    parsed = parse_docx(_make_docx(tmp_path / "c.docx"))
    by_text = {b.text: b.heading for b in parsed.blocks}
    assert by_text["Preambule du contrat."] == "Contrat de prestation"
    assert by_text["Obligations du prestataire."] == "Article 3"


def test_table_rows_are_captured(tmp_path: Path) -> None:
    parsed = parse_docx(_make_docx(tmp_path / "c.docx"))
    joined = " ".join(b.text for b in parsed.blocks)
    assert "Poste" in joined and "Montant" in joined
    assert "Conseil" in joined and "12000" in joined


def test_document_order_is_preserved_across_tables(tmp_path: Path) -> None:
    # python-docx exposes paragraphs and tables separately; reading them in
    # turn would put the table after 'Paragraphe final.'
    parsed = parse_docx(_make_docx(tmp_path / "c.docx"))
    texts = [b.text for b in parsed.blocks]
    table_index = next(i for i, t in enumerate(texts) if "Conseil" in t)
    final_index = next(i for i, t in enumerate(texts) if t == "Paragraphe final.")
    assert table_index < final_index


def test_empty_paragraphs_are_dropped(tmp_path: Path) -> None:
    path = tmp_path / "sparse.docx"
    document = Document()
    document.add_paragraph("Contenu.")
    document.add_paragraph("   ")
    document.add_paragraph("")
    document.save(str(path))
    assert [b.text for b in parse_docx(path).blocks] == ["Contenu."]


def test_french_accents_survive(tmp_path: Path) -> None:
    path = tmp_path / "fr.docx"
    document = Document()
    document.add_paragraph("Résiliation à effet immédiat.")
    document.save(str(path))
    assert parse_docx(path).blocks[0].text == "Résiliation à effet immédiat."


def test_a_corrupt_docx_raises_a_readable_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.docx"
    path.write_bytes(b"not a docx at all")
    with pytest.raises(ValueError) as excinfo:
        parse_docx(path)
    assert "broken.docx" in str(excinfo.value)
