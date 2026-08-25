"""XLSX parsing: sheets, headers, and empty-cell handling."""

from __future__ import annotations

from pathlib import Path

import pytest
from openpyxl import Workbook

from parsers.xlsx import parse_xlsx


def _make_xlsx(path: Path) -> Path:
    workbook = Workbook()
    first = workbook.active
    first.title = "Budget"
    first.append(["Poste", "Montant"])
    first.append(["Conseil", 12000])
    first.append(["Formation", 3500])
    second = workbook.create_sheet("Ventes")
    second.append(["Client", "Total"])
    second.append(["Dupont SARL", 8000])
    workbook.save(str(path))
    return path


def test_doc_type_and_both_sheets_are_parsed(tmp_path: Path) -> None:
    parsed = parse_xlsx(_make_xlsx(tmp_path / "compta.xlsx"))
    assert parsed.doc_type == "xlsx"
    assert {b.sheet for b in parsed.blocks} == {"Budget", "Ventes"}


def test_every_row_carries_its_column_headers(tmp_path: Path) -> None:
    # '12000' alone is unanswerable; 'Montant: 12000' is not.
    parsed = parse_xlsx(_make_xlsx(tmp_path / "compta.xlsx"))
    conseil = next(b for b in parsed.blocks if "Conseil" in b.text)
    assert "Poste: Conseil" in conseil.text
    assert "Montant: 12000" in conseil.text


def test_the_header_row_itself_is_not_emitted_as_data(tmp_path: Path) -> None:
    parsed = parse_xlsx(_make_xlsx(tmp_path / "compta.xlsx"))
    budget = [b.text for b in parsed.blocks if b.sheet == "Budget"]
    assert len(budget) == 2  # two data rows, not three


def test_empty_rows_are_skipped(tmp_path: Path) -> None:
    path = tmp_path / "gappy.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["A", "B"])
    sheet.append([None, None])
    sheet.append(["x", "y"])
    workbook.save(str(path))
    assert len(parse_xlsx(path).blocks) == 1


def test_empty_cells_do_not_produce_none_text(tmp_path: Path) -> None:
    path = tmp_path / "partial.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "S"
    sheet.append(["A", "B"])
    sheet.append(["x", None])
    workbook.save(str(path))
    (block,) = parse_xlsx(path).blocks
    assert "None" not in block.text


def test_an_empty_workbook_yields_no_blocks(tmp_path: Path) -> None:
    path = tmp_path / "empty.xlsx"
    Workbook().save(str(path))
    assert parse_xlsx(path).blocks == []


def test_a_corrupt_xlsx_raises_a_readable_error(tmp_path: Path) -> None:
    path = tmp_path / "broken.xlsx"
    path.write_bytes(b"not a workbook")
    with pytest.raises(ValueError) as excinfo:
        parse_xlsx(path)
    assert "broken.xlsx" in str(excinfo.value)
